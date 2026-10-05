import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Colors
    DARK_BG = RGBColor(15, 23, 42)
    TEXT_WHITE = RGBColor(255, 255, 255)
    ACCENT_BLUE = RGBColor(56, 189, 248)
    TEXT_MUTED = RGBColor(148, 163, 184)
    CARD_BG = RGBColor(30, 41, 59)

    blank_layout = prs.slide_layouts[6]

    slides_data = [
        # Slide 1
        {
            "title": "Design and Implementation of a DBMS for Airline Reservation & Flight Operations",
            "content": "Team Members & Role Allocation:\n\n• Mithila — Streamlit UI Development, Testing Suite & Presentation\n• Sukriti — Database Schema Architecture & ER Diagram\n• Hrishika — Python Backend Logic & MySQL Connector Infrastructure\n\nCourse Project | DBMS",
            "type": "title"
        },
        # Slide 2
        {
            "title": "Problem Statement",
            "content": "Traditional airline operational workflows face significant operational friction:\n\n1. Disconnected Data Silos: Passenger profiles, flight schedules, seat allocations, and payment receipts managed separately.\n2. Data Inconsistency & Redundancy: Risk of overbooking flights and missing transaction verification.\n3. Lack of Real-time Visibility: Inability for flight managers to view live metrics on reservations and seat availability."
        },
        # Slide 3
        {
            "title": "Project Objectives",
            "content": "Key goals of the Airline Management System:\n\n• Centralized Data Repository: Unified MySQL relational database storing all core entities.\n• ACID-Compliant Transactions: Guarantee transaction integrity for flight bookings and payments.\n• User-Friendly Streamlit Dashboard: Intuitive interactive UI built with Python.\n• Robust CRUD Operations: Seamless addition, updating, viewing, and deletion of records."
        },
        # Slide 4
        {
            "title": "Scope & Target Users",
            "content": "System Scope & Key User Roles:\n\n• Airline System Administrators: Fleet configuration (Aircraft, Seats), Airport routing.\n• Reservation Desk Staff: Passenger registration, ticket booking, reservation status updates.\n• Financial Operations Officers: Payment recording (UPI, Card, Net Banking) and revenue reconciliation."
        },
        # Slide 5
        {
            "title": "Entity Relationship (ER) Diagram",
            "content": "Comprehensive ER Diagram conceptualizing entity interactions:\n\n• Entities: Passenger, Airport, Aircraft, Flight, Seat, Reservation, Payment.\n• Key Mappings:\n  - Passenger (1) ──< Booking >── (N) Reservation\n  - Flight (1) ──< Associated With >── (N) Reservation\n  - Reservation (1) ──< Settles >── (1) Payment [UNIQUE Constraint]"
        },
        # Slide 6
        {
            "title": "Database Relational Schema",
            "content": "Core Tables and Schemas:\n\n• Passenger(passenger_id PK, name, email UNIQUE, phone, passport_no UNIQUE)\n• Airport(airport_id PK, airport_code UNIQUE, airport_name, city, country)\n• Aircraft(aircraft_id PK, aircraft_model, capacity CHECK>0)\n• Flight(flight_id PK, flight_number UNIQUE, aircraft_id FK, departure_airport_id FK, arrival_airport_id FK, departure_datetime, arrival_datetime, status, base_fare CHECK>=0)\n• Seat(seat_id PK, aircraft_id FK, seat_number, seat_class, UNIQUE(aircraft_id, seat_number))\n• Reservation(reservation_id PK, passenger_id FK, flight_id FK, booking_date, reservation_status)\n• Payment(payment_id PK, reservation_id FK UNIQUE, amount CHECK>=0, payment_method, payment_status, payment_date)"
        },
        # Slide 7
        {
            "title": "Relationships & Cardinalities",
            "content": "Relational Multiplicities & Business Rules:\n\n• Passenger ➔ Reservation: 1 to N (A passenger can create multiple flight bookings).\n• Flight ➔ Reservation: 1 to N (A flight accommodates multiple passenger reservations).\n• Aircraft ➔ Flight: 1 to N (An aircraft performs multiple scheduled flight legs).\n• Reservation ➔ Payment: 1 to 1 (Enforced via UNIQUE constraint on Payment.reservation_id)."
        },
        # Slide 8
        {
            "title": "Normalization Analysis (3NF)",
            "content": "Verification of Third Normal Form (3NF):\n\n• 1NF: All attributes contain atomic values; composite/multivalued attributes decomposed.\n• 2NF: All non-key attributes fully functionally dependent on primary keys.\n• 3NF: No transitive dependencies exist (e.g. Airport details referenced via Foreign Key airport_id rather than embedded in Flight)."
        },
        # Slide 9
        {
            "title": "SQL Schema Constraints & Integrity Rules",
            "content": "Integrity Rules Enforced in DDL:\n\n• PRIMARY KEY: Auto-incrementing integer IDs across all entities.\n• FOREIGN KEY: Referential integrity enforced across Flight, Reservation, and Payment tables.\n• UNIQUE: Prevents duplicate emails, passport numbers, flight numbers, and duplicate payments for a reservation.\n• CHECK Constraints: base_fare >= 0, capacity > 0, arrival_datetime > departure_datetime."
        },
        # Slide 10
        {
            "title": "Key SQL Queries & Analytics",
            "content": "Implementation of Essential Database Operations:\n\n• Complex Join Queries: Flight schedule retrieval with departure/arrival airport codes.\n• Aggregations: Grouping payments by method (UPI, Card, Net Banking) via SUM() and COUNT().\n• Conditional Filters: Querying flights by price ranking and confirmed booking status."
        },
        # Slide 11
        {
            "title": "Streamlit User Interface (Mithila UI)",
            "content": "Interactive Desktop Web Application Features:\n\n• Modern Glassmorphism Styling: Custom CSS with metric card counters.\n• Modular Navigation Sidebar: Dashboard, Passengers, Flights, Reservations, Payments.\n• Foreign-Key Dropdowns: Intuitive entity selection replacing raw database IDs."
        },
        # Slide 12
        {
            "title": "System Architecture & Connectivity",
            "content": "3-Tier Decoupled Architecture:\n\n┌────────────────────────┐\n│ Streamlit Presentation │  (Mithila UI)\n└───────────┬────────────┘\n            ↓\n┌────────────────────────┐\n│ Python Backend Service │  (Hrishika Backend)\n└───────────┬────────────┘\n            ↓\n┌────────────────────────┐\n│ MySQL Relational DB    │  (Sukriti Database)\n└────────────────────────┘"
        },
        # Slide 13
        {
            "title": "Live Demonstration & Database Verification",
            "content": "End-to-End Data Persistence Verification:\n\n1. UI Passenger Registration ➔ Calls Hrishika Backend ➔ MySQL INSERT ➔ Streamlit Refresh.\n2. MySQL Workbench Verification: SELECT * FROM Passenger reflects new record immediately.\n3. UI Record Deletion ➔ MySQL DELETE executed ➔ UI state updated in real-time."
        },
        # Slide 14
        {
            "title": "Conclusion & Future Enhancements",
            "content": "Project Summary & Future Roadmap:\n\n• Conclusion: Successfully engineered a robust, normalized, 3-tier DBMS for airline operations.\n• Future Enhancements:\n  - Automated Email E-Ticket Generation (PDF Attachment).\n  - Interactive Seat Map Selection UI for passengers.\n  - Role-Based Access Control (RBAC) authentication system."
        }
    ]

    for data in slides_data:
        slide = prs.slides.add_slide(blank_layout)
        
        # Background shape
        bg = slide.shapes.add_shape(1, 0, 0, Inches(13.333), Inches(7.5)) # 1 = msoShapeRectangle
        bg.fill.solid()
        bg.fill.fore_color.rgb = DARK_BG
        bg.line.fill.background()

        # Card container
        card = slide.shapes.add_shape(1, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = ACCENT_BLUE
        card.line.width = Pt(1.5)

        # Title
        tx_box = slide.shapes.add_textbox(Inches(1.2), Inches(1.1), Inches(10.9), Inches(1.0))
        tf = tx_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = data["title"]
        p.font.size = Pt(24 if len(data["title"]) > 40 else 28)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE

        # Content
        content_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.2), Inches(10.9), Inches(4.2))
        ctf = content_box.text_frame
        ctf.word_wrap = True
        cp = ctf.paragraphs[0]
        cp.text = data["content"]
        cp.font.size = Pt(16)
        cp.font.color.rgb = TEXT_WHITE
        
    output_pptx = os.path.join(os.path.dirname(__file__), "project_presentation.pptx")
    prs.save(output_pptx)
    print(f"PowerPoint presentation generated at: {output_pptx}")

if __name__ == "__main__":
    create_presentation()
